from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook


class ExcelDashboardReader:

    NS = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }

    # =========================================================
    # PAGE 1 - TEMPLATE ROWS
    # =========================================================

    ROW_MAPPING = {
        # ACTIVE CUSTOMER DETAILS
        22003: "ios_old",
        22004: "ios_new",

        22006: "web_old",
        22007: "web_new",

        22009: "ksd_old",
        22010: "ksd_new",

        22012: "bus_old",
        22013: "bus_new",

        22015: "reseller_old",
        22016: "reseller_new",

        22018: "family_google",
        22019: "minor_google",

        22021: "invoice_account_1",
        22022: "invoice_account_2",

        # SUMMARY
        22026: "cb_old",
        22027: "cb_new",
        22028: "cb_total",
        22029: "google",
        22030: "invoices",
        22031: "total_subscriptions",
    }

    # =========================================================
    # PAGE 1 - TEMPLATE COLUMNS
    # =========================================================

    COLUMN_MAPPING = {
        "G": "monthly",
        "H": "annual",
        "I": "total",
        "J": "nz_monthly",
        "K": "os_monthly",
        "L": "nz_annual",
        "M": "os_annual",
    }

    # =========================================================
    # PAGE 1 - ACTIVE CUSTOMER DISPLAY INFO
    # =========================================================

    ACTIVE_DETAILS = {
        "ios_old": {
            "active": "iOS",
            "status": "Old",
        },

        "ios_new": {
            "active": "iOS",
            "status": "New",
        },

        "web_old": {
            "active": "Web",
            "status": "Old",
        },

        "web_new": {
            "active": "Web",
            "status": "New",
        },

        "ksd_old": {
            "active": "KSD",
            "status": "Old",
        },

        "ksd_new": {
            "active": "KSD",
            "status": "New",
        },

        "bus_old": {
            "active": "Bus",
            "status": "Old",
        },

        "bus_new": {
            "active": "Bus",
            "status": "New",
        },

        "reseller_old": {
            "active": "Reseller",
            "status": "Old",
        },

        "reseller_new": {
            "active": "Reseller",
            "status": "New",
        },

        "family_google": {
            "active": "Family",
            "status": "Google",
        },

        "minor_google": {
            "active": "Minor",
            "status": "Google",
        },

        "invoice_account_1": {
            "active": "Inv",
            "status": "00 A/c",
        },

        "invoice_account_2": {
            "active": "Inv",
            "status": "00 A/c",
        },
    }

    # =========================================================
    # PAGE 1 - SUMMARY DISPLAY INFO
    # =========================================================

    SUMMARY_DETAILS = {
        "cb_old": {
            "name": "CB OLD",
        },

        "cb_new": {
            "name": "CB NEW",
        },

        "cb_total": {
            "name": "CB Total",
        },

        "google": {
            "name": "GOOGLE",
        },

        "invoices": {
            "name": "INVOICES",
        },

        "total_subscriptions": {
            "name": "Total Subscriptions",
        },
    }

    # =========================================================
    # INITIALISE
    # =========================================================

    def __init__(self, file_path):

        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"Excel file not found: {self.file_path}"
            )

    # =========================================================
    # FIND SHEET
    # =========================================================

    def _find_sheet(self, xlsx, sheet_name):

        workbook_xml = ET.fromstring(
            xlsx.read("xl/workbook.xml")
        )

        relationships_xml = ET.fromstring(
            xlsx.read(
                "xl/_rels/workbook.xml.rels"
            )
        )

        relationships = {}

        for relationship in relationships_xml:

            relationships[
                relationship.attrib["Id"]
            ] = relationship.attrib["Target"]

        sheets = workbook_xml.find(
            "main:sheets",
            self.NS,
        )

        if sheets is None:
            raise RuntimeError(
                "Unable to find sheets in workbook."
            )

        for sheet in sheets:

            if sheet.attrib["name"] != sheet_name:
                continue

            relationship_id = sheet.attrib[
                "{http://schemas.openxmlformats.org/"
                "officeDocument/2006/relationships}id"
            ]

            target = relationships[
                relationship_id
            ]

            target = target.lstrip("/")

            if target.startswith("xl/"):
                return target

            return "xl/" + target

        raise RuntimeError(
            f"Sheet '{sheet_name}' not found."
        )

    # =========================================================
    # PAGE 1 - FIND TEMPLATE
    # =========================================================

    def _find_template_sheet(self, xlsx):

        return self._find_sheet(
            xlsx,
            "Template",
        )

    # =========================================================
    # PAGE 1 - READ TEMPLATE ROWS
    # =========================================================

    def _read_template_rows(self):

        with ZipFile(
            self.file_path,
            "r",
        ) as xlsx:

            template_path = (
                self._find_template_sheet(
                    xlsx
                )
            )

            sheet_xml = ET.fromstring(
                xlsx.read(
                    template_path
                )
            )

            rows = {}

            for row in sheet_xml.findall(
                ".//main:row",
                self.NS,
            ):

                row_number = int(
                    row.attrib["r"]
                )

                if row_number not in self.ROW_MAPPING:
                    continue

                row_data = {}

                for cell in row.findall(
                    "main:c",
                    self.NS,
                ):

                    cell_reference = (
                        cell.attrib["r"]
                    )

                    column = "".join(
                        character
                        for character in cell_reference
                        if character.isalpha()
                    )

                    if column not in self.COLUMN_MAPPING:
                        continue

                    value_element = cell.find(
                        "main:v",
                        self.NS,
                    )

                    if value_element is None:
                        value = 0
                    else:
                        value = value_element.text

                    try:
                        value = int(
                            float(value)
                        )

                    except (
                        TypeError,
                        ValueError,
                    ):
                        value = 0

                    field_name = (
                        self.COLUMN_MAPPING[
                            column
                        ]
                    )

                    row_data[
                        field_name
                    ] = value

                # Ensure every expected field exists
                for field_name in (
                    self.COLUMN_MAPPING.values()
                ):

                    if field_name not in row_data:
                        row_data[
                            field_name
                        ] = 0

                rows[
                    row_number
                ] = row_data

            return rows

    # =========================================================
    # PAGE 1 - EMPTY ROW
    # =========================================================

    def _empty_row(self):

        return {
            field: 0
            for field
            in self.COLUMN_MAPPING.values()
        }

    # =========================================================
    # PAGE 2 - ACTIVE SUBSCRIPTIONS
    # =========================================================
    #
    # Workbook:
    # Active Subs & Graphics.xlsx
    #
    # Sheet:
    # ACTIVE SUBS
    #
    # Actual workbook columns:
    #
    # B = Year
    # C = Month
    # D = Active
    # N = Non-renewing
    # V = Balance Month
    # W = Balance VALUE
    #
    # Data begins at row 7.
    #
    # =========================================================

    def _read_active_subs_sheet(self):

        workbook = load_workbook(
            filename=self.file_path,
            data_only=True,
            read_only=True,
        )

        try:

            if "ACTIVE SUBS" not in workbook.sheetnames:

                raise RuntimeError(
                    "Sheet 'ACTIVE SUBS' not found."
                )

            sheet = workbook[
                "ACTIVE SUBS"
            ]

            active_subs = []

            current_year = None

            # -------------------------------------------------
            # NUMBER CONVERTER
            # -------------------------------------------------

            def to_number(value):

                if value is None:
                    return 0

                try:
                    return int(
                        float(value)
                    )

                except (
                    TypeError,
                    ValueError,
                ):
                    return 0

            # -------------------------------------------------
            # READ MONTHLY DATA
            # -------------------------------------------------

            for row_number in range(
                7,
                sheet.max_row + 1,
            ):

                # ---------------------------------------------
                # READ ACTUAL EXCEL CELLS
                # ---------------------------------------------

                year_value = sheet[
                    f"B{row_number}"
                ].value

                month_value = sheet[
                    f"C{row_number}"
                ].value

                active_value = sheet[
                    f"D{row_number}"
                ].value

                non_renewing_value = sheet[
                    f"N{row_number}"
                ].value

                # IMPORTANT:
                # W contains the actual Balance number.
                balance_value = sheet[
                    f"W{row_number}"
                ].value

                # ---------------------------------------------
                # MONTH
                # ---------------------------------------------

                month = str(
                    month_value or ""
                ).strip()

                # Ignore rows that do not contain
                # a month.
                if not month:
                    continue

                # ---------------------------------------------
                # YEAR
                # ---------------------------------------------

                if year_value is not None:

                    try:
                        current_year = int(
                            float(
                                year_value
                            )
                        )

                    except (
                        TypeError,
                        ValueError,
                    ):
                        pass

                # ---------------------------------------------
                # VALUES
                # ---------------------------------------------

                active = to_number(
                    active_value
                )

                non_renewing = to_number(
                    non_renewing_value
                )

                # Read Balance DIRECTLY from Excel column W.
                balance = to_number(
                    balance_value
                )

                # ---------------------------------------------
                # LABEL
                # ---------------------------------------------

                if current_year is not None:

                    label = (
                        f"{month.title()} "
                        f"{current_year}"
                    )

                else:

                    label = month.title()

                # ---------------------------------------------
                # ADD RECORD
                # ---------------------------------------------

                active_subs.append(
                    {
                        "row": row_number,
                        "year": current_year,
                        "month": month,
                        "label": label,
                        "active": active,
                        "non_renewing": non_renewing,
                        "balance": balance,
                    }
                )

            return active_subs

        finally:

            workbook.close()
        # =========================================================
    # PAGE 3 - TRIALISTS
    # =========================================================
    #
    # Workbook:
    # Trialists Graphics.xlsx
    #
    # Sheet:
    # TRIALISTS
    #
    # Main source table:
    #
    # I = Year
    # J = Month
    # K = Source
    # L = Total Trialists
    # M = Converted This Month
    # N = Converted Next Month 1-7 Days
    # O = No / Not Converted
    #
    # Data begins at row 13.
    #
    # Each month contains:
    # CB iOS
    # CB Web
    # Google
    #
    # =========================================================

    def _read_trialists_sheet(self):

        workbook = load_workbook(
            filename=self.file_path,
            data_only=True,
            read_only=True,
        )

        try:

            if "TRIALISTS" not in workbook.sheetnames:

                raise RuntimeError(
                    "Sheet 'TRIALISTS' not found."
                )

            sheet = workbook["TRIALISTS"]

            # -------------------------------------------------
            # NUMBER CONVERTER
            # -------------------------------------------------

            def to_number(value):

                if value is None:
                    return 0

                try:
                    return int(
                        float(value)
                    )

                except (
                    TypeError,
                    ValueError,
                ):
                    return 0

            details = []

            current_year = None
            current_month = None

            # -------------------------------------------------
            # READ SOURCE-LEVEL DATA
            # -------------------------------------------------

            for row_number in range(
                13,
                sheet.max_row + 1,
            ):

                year_value = sheet[
                    f"I{row_number}"
                ].value

                month_value = sheet[
                    f"J{row_number}"
                ].value

                source_value = sheet[
                    f"K{row_number}"
                ].value

                total_value = sheet[
                    f"L{row_number}"
                ].value

                this_month_value = sheet[
                    f"M{row_number}"
                ].value

                next_7_days_value = sheet[
                    f"N{row_number}"
                ].value

                no_value = sheet[
                    f"O{row_number}"
                ].value

                # ---------------------------------------------
                # SOURCE
                # ---------------------------------------------

                source = str(
                    source_value or ""
                ).strip()

                # Only process the actual trialist source rows.
                if source not in {
                    "CB iOS",
                    "CB Web",
                    "Google",
                }:
                    continue

                # ---------------------------------------------
                # YEAR
                # ---------------------------------------------

                if year_value is not None:

                    try:
                        current_year = int(
                            float(year_value)
                        )

                    except (
                        TypeError,
                        ValueError,
                    ):
                        pass

                # ---------------------------------------------
                # MONTH
                # ---------------------------------------------

                if month_value is not None:

                    month_text = str(
                        month_value
                    ).strip()

                    if month_text:
                        current_month = month_text

                # ---------------------------------------------
                # VALUES
                # ---------------------------------------------

                total = to_number(
                    total_value
                )

                this_month = to_number(
                    this_month_value
                )

                next_7_days = to_number(
                    next_7_days_value
                )

                no_response = to_number(
                    no_value
                )

                # ---------------------------------------------
                # CONVERTED
                # ---------------------------------------------

                converted = (
                    this_month
                    + next_7_days
                )

                # ---------------------------------------------
                # CONVERSION RATE
                # ---------------------------------------------

                if total > 0:

                    conversion_rate = round(
                        (
                            converted
                            / total
                        ) * 100,
                        2,
                    )

                else:

                    conversion_rate = 0

                # ---------------------------------------------
                # LABEL
                # ---------------------------------------------

                if (
                    current_month
                    and current_year
                ):

                    label = (
                        f"{current_month.title()} "
                        f"{current_year}"
                    )

                elif current_month:

                    label = (
                        current_month.title()
                    )

                else:

                    label = ""

                # ---------------------------------------------
                # ADD DETAIL RECORD
                # ---------------------------------------------

                details.append(
                    {
                        "row": row_number,
                        "year": current_year,
                        "month": current_month,
                        "label": label,
                        "source": source,
                        "total": total,
                        "this_month": this_month,
                        "next_7_days": next_7_days,
                        "converted": converted,
                        "no_response": no_response,
                        "conversion_rate":
                            conversion_rate,
                    }
                )

            # =================================================
            # BUILD MONTHLY SUMMARY
            # =================================================

            monthly_map = {}

            for item in details:

                key = (
                    item["year"],
                    item["month"],
                )

                if key not in monthly_map:

                    monthly_map[key] = {
                        "year":
                            item["year"],

                        "month":
                            item["month"],

                        "label":
                            item["label"],

                        "total":
                            0,

                        "this_month":
                            0,

                        "next_7_days":
                            0,

                        "converted":
                            0,

                        "no_response":
                            0,
                    }

                monthly_map[
                    key
                ]["total"] += (
                    item["total"]
                )

                monthly_map[
                    key
                ]["this_month"] += (
                    item["this_month"]
                )

                monthly_map[
                    key
                ]["next_7_days"] += (
                    item["next_7_days"]
                )

                monthly_map[
                    key
                ]["converted"] += (
                    item["converted"]
                )

                monthly_map[
                    key
                ]["no_response"] += (
                    item["no_response"]
                )

            monthly = list(
                monthly_map.values()
            )

            # -------------------------------------------------
            # MONTHLY CONVERSION RATE
            # -------------------------------------------------

            for item in monthly:

                if item["total"] > 0:

                    item[
                        "conversion_rate"
                    ] = round(
                        (
                            item["converted"]
                            / item["total"]
                        ) * 100,
                        2,
                    )

                else:

                    item[
                        "conversion_rate"
                    ] = 0

            # =================================================
            # SOURCE SUMMARY
            # =================================================

            source_map = {}

            for item in details:

                source = item[
                    "source"
                ]

                if source not in source_map:

                    source_map[source] = {
                        "source":
                            source,

                        "total":
                            0,

                        "this_month":
                            0,

                        "next_7_days":
                            0,

                        "converted":
                            0,

                        "no_response":
                            0,
                    }

                source_map[
                    source
                ]["total"] += (
                    item["total"]
                )

                source_map[
                    source
                ]["this_month"] += (
                    item["this_month"]
                )

                source_map[
                    source
                ]["next_7_days"] += (
                    item["next_7_days"]
                )

                source_map[
                    source
                ]["converted"] += (
                    item["converted"]
                )

                source_map[
                    source
                ]["no_response"] += (
                    item["no_response"]
                )

            sources = list(
                source_map.values()
            )

            for item in sources:

                if item["total"] > 0:

                    item[
                        "conversion_rate"
                    ] = round(
                        (
                            item["converted"]
                            / item["total"]
                        ) * 100,
                        2,
                    )

                else:

                    item[
                        "conversion_rate"
                    ] = 0

            # =================================================
            # LATEST MONTH
            # =================================================

            latest = (
                monthly[-1]
                if monthly
                else None
            )

            # =================================================
            # RETURN PAGE 3 DATA
            # =================================================

            return {
                "details": details,
                "monthly": monthly,
                "sources": sources,
                "latest": latest,
            }

        finally:

            workbook.close()

    # =========================================================
    # INTERNATIONAL — FF INTERNATIONAL WORKBOOK
    # =========================================================
    #
    # Workbook:
    # FF - International - Final Dashboard tables.xlsx
    #
    # Sheet: Income Source $
    #
    # B = Month
    # C = Country
    # D = Licencing
    # E = Marketing Support
    # F = Other
    # G = Product Sales
    # H = Shipping
    # I = Apple
    # J = Google
    # K = Website (Stripe)
    # L = Invoices
    # M = Resellers
    # N = Total
    #
    # Headings on Excel row 4. Data begins at row 6.
    #
    # Sheet: Top Countries $
    #
    # A = Month
    # B = Order
    # C = Country
    # D = Total
    #
    # Headings on Excel row 6. Data begins at row 8.
    #
    # =========================================================

    INCOME_SOURCE_COLUMNS = {
        "D": "licensing",
        "E": "marketing_support",
        "F": "other",
        "G": "product_sales",
        "H": "shipping",
        "I": "apple",
        "J": "google",
        "K": "website_stripe",
        "L": "invoices",
        "M": "resellers",
        "N": "total",
    }

    def _format_month_label(self, month_value):

        if month_value is None:
            return ""

        if hasattr(month_value, "strftime"):
            return month_value.strftime("%b %Y")

        try:
            parsed = datetime.fromisoformat(
                str(month_value).strip()
            )

            return parsed.strftime("%b %Y")

        except ValueError:
            return str(month_value).strip()

    def _read_international_workbook(self):

        workbook = load_workbook(
            filename=self.file_path,
            data_only=True,
            read_only=True,
        )

        try:

            if "Income Source $" not in workbook.sheetnames:
                raise RuntimeError(
                    "Sheet 'Income Source $' not found."
                )

            if "Top Countries $" not in workbook.sheetnames:
                raise RuntimeError(
                    "Sheet 'Top Countries $' not found."
                )

            income_sheet = workbook["Income Source $"]
            countries_sheet = workbook["Top Countries $"]

            def to_number(value):

                if value is None:
                    return 0

                try:
                    return float(value)

                except (
                    TypeError,
                    ValueError,
                ):
                    return 0

            # -------------------------------------------------
            # INCOME SOURCE $
            # -------------------------------------------------

            income_monthly = []

            for row_number in range(
                6,
                income_sheet.max_row + 1,
            ):

                month_value = income_sheet[
                    f"B{row_number}"
                ].value

                country_value = income_sheet[
                    f"C{row_number}"
                ].value

                if month_value is None:
                    continue

                month_label = self._format_month_label(
                    month_value
                )

                record = {
                    "row": row_number,
                    "month": month_value,
                    "label": month_label,
                    "country": str(
                        country_value or ""
                    ).strip(),
                }

                for (
                    column,
                    field_name,
                ) in self.INCOME_SOURCE_COLUMNS.items():

                    record[
                        field_name
                    ] = to_number(
                        income_sheet[
                            f"{column}{row_number}"
                        ].value
                    )

                income_monthly.append(record)

            income_sources = []

            for (
                column,
                field_name,
            ) in self.INCOME_SOURCE_COLUMNS.items():

                if field_name == "total":
                    continue

                total_value = sum(
                    item[field_name]
                    for item in income_monthly
                )

                income_sources.append(
                    {
                        "source": field_name,
                        "total": total_value,
                    }
                )

            income_sources.sort(
                key=lambda item: item["total"],
                reverse=True,
            )

            income_latest = (
                income_monthly[-1]
                if income_monthly
                else None
            )

            # -------------------------------------------------
            # TOP COUNTRIES $
            # -------------------------------------------------

            country_details = []

            for row_number in range(
                8,
                countries_sheet.max_row + 1,
            ):

                month_value = countries_sheet[
                    f"A{row_number}"
                ].value

                order_value = countries_sheet[
                    f"B{row_number}"
                ].value

                country_value = countries_sheet[
                    f"C{row_number}"
                ].value

                total_value = countries_sheet[
                    f"D{row_number}"
                ].value

                if month_value is None:
                    continue

                country_name = str(
                    country_value or ""
                ).strip()

                if not country_name:
                    continue

                month_label = self._format_month_label(
                    month_value
                )

                country_details.append(
                    {
                        "row": row_number,
                        "month": month_value,
                        "label": month_label,
                        "order": (
                            int(float(order_value))
                            if order_value is not None
                            and str(order_value).strip()
                            else None
                        ),
                        "country": country_name,
                        "total": to_number(
                            total_value
                        ),
                        "is_total_row":
                            country_name.lower()
                            == "total",
                    }
                )

            country_monthly_map = {}

            for item in country_details:

                if item["is_total_row"]:
                    continue

                key = item["label"]

                if key not in country_monthly_map:

                    country_monthly_map[key] = {
                        "label": item["label"],
                        "month": item["month"],
                        "countries": [],
                        "total": 0,
                    }

                country_monthly_map[
                    key
                ]["countries"].append(
                    {
                        "order":
                            item["order"],
                        "country":
                            item["country"],
                        "total":
                            item["total"],
                    }
                )

                country_monthly_map[
                    key
                ]["total"] += item["total"]

            country_monthly = []

            for item in income_monthly:

                label = item["label"]

                month_countries = (
                    country_monthly_map.get(
                        label,
                        {
                            "label": label,
                            "month": item["month"],
                            "countries": [],
                            "total": 0,
                        },
                    )
                )

                month_countries[
                    "countries"
                ].sort(
                    key=lambda country: (
                        country["order"]
                        if country["order"]
                        is not None
                        else 999,
                        -country["total"],
                    )
                )

                country_monthly.append(
                    month_countries
                )

            country_totals_map = {}

            for item in country_details:

                if item["is_total_row"]:
                    continue

                country_name = item["country"]

                if country_name not in (
                    country_totals_map
                ):

                    country_totals_map[
                        country_name
                    ] = {
                        "country":
                            country_name,
                        "total": 0,
                    }

                country_totals_map[
                    country_name
                ]["total"] += item["total"]

            country_totals = list(
                country_totals_map.values()
            )

            country_totals.sort(
                key=lambda item: item["total"],
                reverse=True,
            )

            countries_latest = (
                country_monthly[-1]
                if country_monthly
                else None
            )

            return {
                "income": {
                    "monthly":
                        income_monthly,
                    "sources":
                        income_sources,
                    "latest":
                        income_latest,
                },
                "countries": {
                    "details":
                        country_details,
                    "monthly":
                        country_monthly,
                    "totals":
                        country_totals,
                    "latest":
                        countries_latest,
                },
            }

        finally:

            workbook.close()

    # =========================================================
    # PAGE 1 - READ DASHBOARD
    # =========================================================

    def read_dashboard(self):

        rows = self._read_template_rows()

        dashboard = {}

        # =====================================================
        # ACTIVE CUSTOMER DETAILS
        # =====================================================

        active = []

        for (
            row_number,
            name,
        ) in self.ROW_MAPPING.items():

            if name not in self.ACTIVE_DETAILS:
                continue

            values = rows.get(
                row_number,
                self._empty_row(),
            )

            details = (
                self.ACTIVE_DETAILS[
                    name
                ]
            )

            active.append(
                {
                    "row": row_number,

                    "active":
                        details["active"],

                    "status":
                        details["status"],

                    "monthly":
                        values["monthly"],

                    "annual":
                        values["annual"],

                    "total":
                        values["total"],

                    "nz_monthly":
                        values["nz_monthly"],

                    "os_monthly":
                        values["os_monthly"],

                    "nz_annual":
                        values["nz_annual"],

                    "os_annual":
                        values["os_annual"],
                }
            )

        dashboard[
            "active_details"
        ] = active

        # =====================================================
        # SUMMARY
        # =====================================================

        summary = {}

        for (
            row_number,
            name,
        ) in self.ROW_MAPPING.items():

            if name not in self.SUMMARY_DETAILS:
                continue

            values = rows.get(
                row_number,
                self._empty_row(),
            )

            summary[
                name
            ] = {
                "row":
                    row_number,

                "name":
                    self.SUMMARY_DETAILS[
                        name
                    ]["name"],

                "monthly":
                    values["monthly"],

                "annual":
                    values["annual"],

                "total":
                    values["total"],

                "nz_monthly":
                    values["nz_monthly"],

                "os_monthly":
                    values["os_monthly"],

                "nz_annual":
                    values["nz_annual"],

                "os_annual":
                    values["os_annual"],
            }

        dashboard[
            "summary"
        ] = summary

        # =====================================================
        # BACKWARD COMPATIBILITY
        # =====================================================
        #
        # Keeps your existing Page 1 React dashboard working.
        #
        # =====================================================

        for name in self.SUMMARY_DETAILS:

            dashboard[
                name
            ] = summary[
                name
            ]

        return dashboard


# =============================================================
# LOCAL TEST
# =============================================================

if __name__ == "__main__":

    project_root = (
        Path(__file__).resolve().parents[2]
    )

    # =========================================================
    # PAGE 1 TEST
    # =========================================================

    customer_file = (
        project_root
        / "Active Customer #.xlsx"
    )

    customer_reader = (
        ExcelDashboardReader(
            customer_file
        )
    )

    customer_dashboard = (
        customer_reader.read_dashboard()
    )

    print(
        "\n===================================="
    )

    print(
        "ACTIVE CUSTOMER DETAILS"
    )

    print(
        "===================================="
    )

    for item in customer_dashboard[
        "active_details"
    ]:

        print(
            f'\nRow {item["row"]} | '
            f'{item["active"]} | '
            f'{item["status"]}'
        )

        print(
            f'  Monthly:      '
            f'{item["monthly"]}'
        )

        print(
            f'  Annual:       '
            f'{item["annual"]}'
        )

        print(
            f'  Total:        '
            f'{item["total"]}'
        )

        print(
            f'  NZ Monthly:   '
            f'{item["nz_monthly"]}'
        )

        print(
            f'  OS Monthly:   '
            f'{item["os_monthly"]}'
        )

        print(
            f'  NZ Annual:    '
            f'{item["nz_annual"]}'
        )

        print(
            f'  OS Annual:    '
            f'{item["os_annual"]}'
        )

    # =========================================================
    # PAGE 1 SUMMARY
    # =========================================================

    print(
        "\n===================================="
    )

    print(
        "SUMMARY"
    )

    print(
        "===================================="
    )

    for (
        name,
        values,
    ) in customer_dashboard[
        "summary"
    ].items():

        print(
            f'\n{values["name"]}'
        )

        print(
            f'  Monthly:      '
            f'{values["monthly"]}'
        )

        print(
            f'  Annual:       '
            f'{values["annual"]}'
        )

        print(
            f'  Total:        '
            f'{values["total"]}'
        )

        print(
            f'  NZ Monthly:   '
            f'{values["nz_monthly"]}'
        )

        print(
            f'  OS Monthly:   '
            f'{values["os_monthly"]}'
        )

        print(
            f'  NZ Annual:    '
            f'{values["nz_annual"]}'
        )

        print(
            f'  OS Annual:    '
            f'{values["os_annual"]}'
        )

    # =========================================================
    # PAGE 2 TEST
    # =========================================================

    active_subs_file = (
        project_root
        / "Active Subs & Graphics.xlsx"
    )

    active_subs_reader = (
        ExcelDashboardReader(
            active_subs_file
        )
    )

    active_subs = (
        active_subs_reader
        ._read_active_subs_sheet()
    )

    print(
        "\n===================================="
    )

    print(
        "ACTIVE SUBS"
    )

    print(
        "===================================="
    )

    for item in active_subs:

        print(
            f'\nRow {item["row"]} | '
            f'{item["label"]}'
        )

        print(
            f'  Active:        '
            f'{item["active"]}'
        )

        print(
            f'  Non-renewing:  '
            f'{item["non_renewing"]}'
        )

        print(
            f'  Balance:       '
            f'{item["balance"]}'
        )

    print(
        "\n===================================="
    )

    print(
        "FINISHED"
    )

    print(
        "===================================="
    )
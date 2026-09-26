from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import re
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook


MISSING_TEXT = {"", "n/a", "na", "none", "null", "-", "—"}


@dataclass
class MetricSeries:
    key: str
    label: str
    category: str
    unit: str
    source: str
    values: pd.Series
    contains_fictitious_values: bool = False
    experimental: bool = True
    notes: list[str] = field(default_factory=list)
    contributions: dict[str, pd.Series] = field(default_factory=dict)
    related: dict[str, Any] = field(default_factory=dict)

    def validation(self) -> dict[str, Any]:
        series = self.values.sort_index()
        usable = series.dropna()
        if usable.empty:
            return {
                "available": False,
                "reason": "No numeric historical observations were found.",
                "observations": 0,
                "missing_months": [],
                "zero_observations": 0,
            }

        expected = pd.date_range(usable.index.min(), usable.index.max(), freq="MS")
        aligned = series.reindex(expected)
        missing = [date.strftime("%Y-%m") for date in aligned[aligned.isna()].index]
        observations = int(usable.shape[0])
        zero_observations = int((usable == 0).sum())
        unique_values = int(usable.nunique())
        nonzero_observations = int((usable != 0).sum())

        reason = None
        if observations < 10:
            reason = "Fewer than 10 usable monthly observations."
        elif len(missing) > max(2, int(len(expected) * 0.2)):
            reason = "Too many gaps in the monthly history."
        elif unique_values < 3:
            reason = "Insufficient historical variation."
        elif nonzero_observations < 4:
            reason = "Too few non-zero historical observations."

        return {
            "available": reason is None,
            "reason": reason,
            "observations": observations,
            "first_actual": usable.index.min().strftime("%Y-%m"),
            "last_actual": usable.index.max().strftime("%Y-%m"),
            "missing_months": missing,
            "zero_observations": zero_observations,
            "unique_values": unique_values,
        }


class ForecastDataLoader:
    """Loads forecasting workbooks without changing the dashboard readers."""

    MRR_FILE_HINT = "ff mrr chargebee google apple stripe"
    TRIALIST_FILE_HINT = "ff trialists total"
    INTERNATIONAL_FILE_HINT = "ff international final dashboard tables"

    def __init__(self, project_root: Path):
        self.project_root = Path(project_root)
        self.source_files = {
            "mrr": self._find_workbook(self.MRR_FILE_HINT),
            "trialists": self._find_workbook(self.TRIALIST_FILE_HINT),
            "international": self._find_workbook(self.INTERNATIONAL_FILE_HINT),
        }
        self.data_quality_issues: list[str] = []

    @staticmethod
    def _normalise_name(value: str) -> str:
        return " ".join(re.findall(r"[a-z0-9]+", value.lower()))

    def _find_workbook(self, hint: str) -> Path:
        hint_tokens = set(hint.split())
        candidates = []
        for path in self.project_root.rglob("*.xlsx"):
            if path.name.startswith("~$"):
                continue
            normalised = self._normalise_name(path.stem)
            score = len(hint_tokens.intersection(normalised.split()))
            if score:
                candidates.append((score, len(normalised), path))

        if not candidates:
            raise FileNotFoundError(f"Unable to locate forecasting workbook matching: {hint}")

        candidates.sort(key=lambda item: (-item[0], item[1], str(item[2])))
        best_score, _, best_path = candidates[0]
        if best_score < max(2, len(hint_tokens) // 2):
            raise FileNotFoundError(f"No reliable workbook match for: {hint}")
        return best_path

    @staticmethod
    def _number(value: Any) -> float | None:
        if value is None or isinstance(value, bool):
            return None
        if isinstance(value, (int, float, np.number)):
            number = float(value)
            return number if np.isfinite(number) else None
        text = str(value).strip()
        if text.lower() in MISSING_TEXT:
            return None
        cleaned = text.replace(",", "").replace("$", "").replace("%", "")
        try:
            number = float(cleaned)
            return number if np.isfinite(number) else None
        except ValueError:
            return None

    @staticmethod
    def _month(value: Any) -> pd.Timestamp | None:
        if value is None:
            return None
        if isinstance(value, (datetime, pd.Timestamp)):
            return pd.Timestamp(value.year, value.month, 1)
        text = str(value).strip().replace("Sept", "Sep")
        if not text:
            return None
        for pattern in ("%b-%y", "%B-%y", "%b %Y", "%B %Y"):
            try:
                parsed_date = datetime.strptime(text, pattern)
                return pd.Timestamp(parsed_date.year, parsed_date.month, 1)
            except ValueError:
                pass
        parsed = pd.to_datetime(text, errors="coerce")
        if pd.isna(parsed):
            return None
        return pd.Timestamp(parsed.year, parsed.month, 1)

    @staticmethod
    def _trim(series: pd.Series) -> pd.Series:
        series = series.sort_index()
        valid = series.dropna()
        if valid.empty:
            return series.iloc[0:0]
        return series.loc[valid.index.min(): valid.index.max()]

    @staticmethod
    def _series(dates: list[pd.Timestamp], values: list[float | None]) -> pd.Series:
        data = pd.Series(values, index=pd.DatetimeIndex(dates), dtype="float64")
        data = data.groupby(level=0).last().sort_index()
        return ForecastDataLoader._trim(data)

    def modification_signature(self) -> tuple[tuple[str, int, int], ...]:
        signature = []
        for key, path in sorted(self.source_files.items()):
            stat = path.stat()
            signature.append((key, stat.st_mtime_ns, stat.st_size))
        return tuple(signature)

    def load(self) -> dict[str, MetricSeries]:
        metrics: dict[str, MetricSeries] = {}
        metrics.update(self._load_mrr_workbook())
        metrics.update(self._load_trialists_workbook())
        metrics.update(self._load_international_workbook())
        return metrics

    def _load_mrr_workbook(self) -> dict[str, MetricSeries]:
        path = self.source_files["mrr"]
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            if "Dashboard Info" not in workbook.sheetnames:
                raise RuntimeError("Sheet 'Dashboard Info' not found in MRR workbook.")
            sheet = workbook["Dashboard Info"]

            header_row = None
            for row in range(1, min(sheet.max_row, 12) + 1):
                values = [str(sheet.cell(row, col).value or "").strip().lower() for col in range(1, 5)]
                if "gateway" in values and any("jan-" in item for item in values):
                    header_row = row
                    break
            if header_row is None:
                header_row = 3

            month_columns: list[tuple[int, pd.Timestamp]] = []
            for column in range(4, sheet.max_column + 1):
                month = self._month(sheet.cell(header_row, column).value)
                if month is not None:
                    month_columns.append((column, month))

            raw: dict[tuple[str, str], pd.Series] = {}
            for row in range(header_row + 1, sheet.max_row + 1):
                gateway = str(sheet.cell(row, 2).value or "").strip()
                metric = str(sheet.cell(row, 3).value or "").strip()
                if not gateway or not metric:
                    continue
                dates = [month for _, month in month_columns]
                values = [self._number(sheet.cell(row, column).value) for column, _ in month_columns]
                raw[(gateway.lower(), metric.lower())] = self._series(dates, values)
        finally:
            workbook.close()

        fictitious_note = (
            "The workbook states that numbers in pink are fictitious. Cell fills were inspected, "
            "but the colours cover broad gateway blocks and cannot reliably identify individual "
            "production versus fictitious observations."
        )

        gateway_names = ["ChargeBee", "Google", "Apple", "Stripe"]
        definitions = {
            "total_revenue": ("Total Revenue", "Financial", "k", "Total Payments - Revenue (k)"),
            "total_mrr": ("Total MRR", "Financial", "k", "Total MRR (k)"),
            "total_active_subscriptions": (
                "Total Active Subscriptions", "Subscriptions", "count", "Total Active Subscriptions"
            ),
            "total_activations": ("Total Activations", "Subscriptions", "count", "Total Activations (New)"),
            "total_signups": ("Total Signups", "Subscriptions", "count", "Total Signups"),
        }

        metrics: dict[str, MetricSeries] = {}
        for key, (label, category, unit, row_label) in definitions.items():
            contributions: dict[str, pd.Series] = {}
            for gateway in gateway_names:
                series = raw.get((gateway.lower(), row_label.lower()))
                if series is not None and not series.empty:
                    contributions[gateway] = series

            if contributions:
                frame = pd.concat(contributions, axis=1).sort_index()
                available_count = frame.notna().sum(axis=1)
                aggregate = frame.sum(axis=1, min_count=1)
                incomplete_zero = (available_count <= 1) & (aggregate == 0)
                aggregate.loc[incomplete_zero] = np.nan
                aggregate = self._trim(aggregate)
            else:
                aggregate = pd.Series(dtype="float64")

            metrics[key] = MetricSeries(
                key=key,
                label=label,
                category=category,
                unit=unit,
                source=path.name,
                values=aggregate,
                contains_fictitious_values=True,
                notes=[fictitious_note, "Aggregate sums only numeric gateway values available for each month."],
                contributions=contributions,
            )

        for gateway in gateway_names:
            slug = gateway.lower()
            for suffix, label_suffix, row_label in (
                ("revenue", "Revenue", "Total Payments - Revenue (k)"),
                ("mrr", "MRR", "Total MRR (k)"),
            ):
                series = raw.get((slug, row_label.lower()), pd.Series(dtype="float64"))
                key = f"{slug}_{suffix}"
                metrics[key] = MetricSeries(
                    key=key,
                    label=f"{gateway} {label_suffix}",
                    category="Gateway Analysis",
                    unit="k",
                    source=path.name,
                    values=self._trim(series),
                    contains_fictitious_values=True,
                    notes=[fictitious_note],
                )

        return metrics

    def _load_trialists_workbook(self) -> dict[str, MetricSeries]:
        path = self.source_files["trialists"]
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            if "TRIALISTS TOTAL" not in workbook.sheetnames:
                raise RuntimeError("Sheet 'TRIALISTS TOTAL' not found.")
            sheet = workbook["TRIALISTS TOTAL"]
            header_row = None
            for row in range(1, min(sheet.max_row, 15) + 1):
                row_values = [str(sheet.cell(row, col).value or "").strip().upper() for col in range(1, 8)]
                if {"YEAR", "MONTH", "TOTAL"}.issubset(set(row_values)):
                    header_row = row
                    break
            if header_row is None:
                raise RuntimeError("Unable to find YEAR/MONTH/TOTAL headers in trialists workbook.")

            header_map = {}
            for column in range(1, sheet.max_column + 1):
                value = str(sheet.cell(header_row, column).value or "").strip().upper()
                if value in {"YEAR", "MONTH", "TOTAL"}:
                    header_map[value] = column

            current_year = None
            dates: list[pd.Timestamp] = []
            values: list[float | None] = []
            for row in range(header_row + 1, sheet.max_row + 1):
                year_value = self._number(sheet.cell(row, header_map["YEAR"]).value)
                if year_value is not None:
                    current_year = int(year_value)
                month_value = sheet.cell(row, header_map["MONTH"]).value
                total_value = self._number(sheet.cell(row, header_map["TOTAL"]).value)
                if current_year is None or month_value is None:
                    continue
                month = self._month(f"{str(month_value).strip()} {current_year}")
                if month is None:
                    continue
                dates.append(month)
                values.append(total_value)
        finally:
            workbook.close()

        return {
            "total_trialists": MetricSeries(
                key="total_trialists",
                label="Total Trialists",
                category="Customer Acquisition",
                unit="count",
                source=path.name,
                values=self._series(dates, values),
                notes=["YEAR is forward-filled; TOTAL is never forward-filled."],
            )
        }

    def _load_international_workbook(self) -> dict[str, MetricSeries]:
        path = self.source_files["international"]
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            if "Income Source $" not in workbook.sheetnames or "Top Countries $" not in workbook.sheetnames:
                raise RuntimeError("International workbook is missing a required sheet.")

            income_sheet = workbook["Income Source $"]
            header_row = None
            for row in range(1, min(15, income_sheet.max_row) + 1):
                values = [str(income_sheet.cell(row, col).value or "").strip() for col in range(1, income_sheet.max_column + 1)]
                if "Month" in values and "Country" in values and "Total" in values:
                    header_row = row
                    break
            if header_row is None:
                raise RuntimeError("Unable to locate international income headers.")

            headers = {
                str(income_sheet.cell(header_row, col).value or "").strip(): col
                for col in range(1, income_sheet.max_column + 1)
                if income_sheet.cell(header_row, col).value is not None
            }
            source_columns = [
                "Licencing", "Marketing Support", "Other", "Product Sales", "Shipping",
                "Apple", "Google", "Website (Stripe)", "Invoices", "Resellers",
            ]
            income_dates: list[pd.Timestamp] = []
            income_values: dict[str, list[float | None]] = {name: [] for name in ["Total", *source_columns]}
            for row in range(header_row + 1, income_sheet.max_row + 1):
                country = str(income_sheet.cell(row, headers["Country"]).value or "").strip()
                if country.lower() != "total":
                    continue
                month = self._month(income_sheet.cell(row, headers["Month"]).value)
                if month is None:
                    continue
                income_dates.append(month)
                for name in income_values:
                    income_values[name].append(self._number(income_sheet.cell(row, headers[name]).value))

            country_sheet = workbook["Top Countries $"]
            country_header = None
            for row in range(1, min(15, country_sheet.max_row) + 1):
                values = [str(country_sheet.cell(row, col).value or "").strip() for col in range(1, 5)]
                if {"Month", "Country", "Total"}.issubset(set(values)):
                    country_header = row
                    break
            if country_header is None:
                raise RuntimeError("Unable to locate international country headers.")
            country_headers = {
                str(country_sheet.cell(country_header, col).value or "").strip(): col
                for col in range(1, country_sheet.max_column + 1)
                if country_sheet.cell(country_header, col).value is not None
            }
            country_records = []
            excluded_country_labels = set()
            for row in range(country_header + 1, country_sheet.max_row + 1):
                country = str(country_sheet.cell(row, country_headers["Country"]).value or "").strip()
                month = self._month(country_sheet.cell(row, country_headers["Month"]).value)
                value = self._number(country_sheet.cell(row, country_headers["Total"]).value)
                if not country or country.lower() == "total" or month is None or value is None:
                    continue
                if country.lower() == "other" or re.fullmatch(r"[a-z]{3}", country):
                    excluded_country_labels.add(country)
                    continue
                country_records.append((month, country, value))
        finally:
            workbook.close()

        if excluded_country_labels:
            self.data_quality_issues.append(
                "Country analysis excluded non-country labels: " + ", ".join(sorted(excluded_country_labels))
            )

        fictitious_note = "The workbook explicitly states that the table values are fictitious."
        metrics: dict[str, MetricSeries] = {}
        source_series = {
            name: self._series(income_dates, values)
            for name, values in income_values.items()
        }
        metrics["international_total_income"] = MetricSeries(
            key="international_total_income",
            label="International Total Income",
            category="Financial",
            unit="currency",
            source=path.name,
            values=source_series["Total"],
            contains_fictitious_values=True,
            notes=[fictitious_note],
            contributions={name: source_series[name] for name in source_columns},
        )

        for name in source_columns:
            key = "international_" + re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
            metrics[key] = MetricSeries(
                key=key,
                label=f"International {name}",
                category="International",
                unit="currency",
                source=path.name,
                values=source_series[name],
                contains_fictitious_values=True,
                notes=[fictitious_note],
            )

        if country_records:
            country_frame = pd.DataFrame(country_records, columns=["month", "country", "value"])
            totals = country_frame.groupby("country")["value"].sum().sort_values(ascending=False)
            top_countries = []
            for country in totals.head(5).index:
                rows = country_frame[country_frame["country"] == country].sort_values("month")
                series = self._series(rows["month"].tolist(), rows["value"].tolist())
                key = "country_" + re.sub(r"[^a-z0-9]+", "_", country.lower()).strip("_")
                metric = MetricSeries(
                    key=key,
                    label=country,
                    category="International — Top Countries",
                    unit="currency",
                    source=path.name,
                    values=series,
                    contains_fictitious_values=True,
                    notes=[fictitious_note],
                    related={"historical_total": float(totals[country])},
                )
                metrics[key] = metric
                top_countries.append({
                    "key": key,
                    "country": country,
                    "historical_total": float(totals[country]),
                })
            metrics["international_total_income"].related["top_countries"] = top_countries

        return metrics
